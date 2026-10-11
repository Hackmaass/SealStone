"""Builds the Sealstone pitch deck on the Vortexa 3.0 template. Re-runnable: `python build.py` -> Sealstone_Vortexa3_Deck.pptx.
Screenshots are read from docs/pitch/img and web/public/landing, so retaking them and rebuilding refreshes the deck."""
import os
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from kit import *  # noqa: F401,F403

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = r"D:\Projects\Vortexa_3.0"
DECK_IMG = os.path.join(REPO, "docs", "pitch", "img")
LAND_IMG = os.path.join(REPO, "web", "public", "landing")
SEALS = os.path.join(HERE, "seals")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "Sealstone_Vortexa3_Deck.pptx")

H_V1 = "9d725c361da710d6bd70c3e2e41b0d2869ad89168c1fb2f864b688f89e4d7146"  # the file the demo vault is waiting on
H_V2 = "c0a1556780624c77f1123936ab3aa445e18c3b2e7b0f3f8163c5986be6ac04db"  # the look-alike
H_BRAND = "18d6890f91dc5dc428f22e2cb402f41c614e0ec0bce44b827238ad9d0f2d44cf"  # keccak256("sealstone")
seal_png(os.path.join(SEALS, "brand.png"), H_BRAND, NAVY)
seal_png(os.path.join(SEALS, "brand_coral.png"), H_BRAND, CORAL)
seal_png(os.path.join(SEALS, "v1.png"), H_V1, NAVY)
seal_png(os.path.join(SEALS, "v2.png"), H_V2, RED)

prs = Presentation(os.path.join(HERE, "struct.pptx"))
S = list(prs.slides)
s_title, s_prob, s_idea, s_prod, s_t1, s_t2, s_t3, s_feas, s_imp, s_diff, s_thanks = S


def seal(slide, name, x, y, d):
    return slide.shapes.add_picture(os.path.join(SEALS, name), I(x), I(y), width=I(d), height=I(d))


# ───────────────────────── 1. title ─────────────────────────
def edit_runs(slide, shape_name, value, size, color=INK, font=HR, width=None, height=None):
    sh = next(s for s in slide.shapes if s.name == shape_name)
    p = sh.text_frame.paragraphs[0]
    runs = p.runs
    if shape_name == "TextBox 5":
        runs[2].text = value
    else:
        r = p.add_run()
        r.text = " " + value
        r.font.name, r.font.size = font, Pt(size)
        r.font.color.rgb = rgb(color)
    if width:
        sh.width = I(width)
    if height:
        sh.height = I(height)
    return sh


edit_runs(s_title, "TextBox 5", "3, B-03", 51.48)
edit_runs(s_title, "TextBox 4", "MILESTONEVAULT: ADVERSARIAL SMART-CONTRACT ESCROW", 51.48, width=15.5, height=1.9)
LEFT = 0.81
rows_t = [("TextBox 3", "TORVALDS", 6.85), ("TextBox 6", "XQR1520246", 7.8)]
for nm, val, y in rows_t:
    sh = edit_runs(s_title, nm, val, 47.71 if nm == "TextBox 3" else 44.72, width=14, height=0.85)
    sh.top = I(y)
    sh.text_frame.word_wrap = False
import copy
for label, val, y in [("TEAM LEADER :", "Omkar Rane", 8.75), ("TEAM MEMBERS :", "Omkar Rane, Pranay Patil", 9.7)]:
    src = next(s for s in s_title.shapes if s.name == "TextBox 6")
    el = copy.deepcopy(src._element)
    s_title.shapes._spTree.append(el)
    sh = s_title.shapes[-1]
    sh.name = "Team " + label.split()[1].title()
    sh.top = I(y)
    sh.text_frame.paragraphs[0].runs[0].text = label
    for r in sh.text_frame.paragraphs[0].runs[1:]:
        r._r.getparent().remove(r._r)
    r = sh.text_frame.paragraphs[0].add_run()
    r.text = " " + val
    r.font.name, r.font.size = HR, Pt(44.72)
    r.font.color.rgb = rgb(INK)
seal(s_title, "brand.png", 11.0, 6.9, 1.5)
tb(s_title, 12.75, 6.85, 5, 0.9, [[("SEALSTONE", {"bold": True})]], size=56, color=NAVY)
tb(s_title, 12.75, 7.75, 5, 0.6, "Escrow that can't be tricked", size=26, color=PURPLE)
notes(s_title, "0:00 (20 s). We are Sealstone, our answer to B-03 MilestoneVault. One sentence: escrow that can't be tricked by ambiguous agreements, swapped deliverables, or malicious contracts. Everything you will see today is a working product you can test.")

# ───────────────────────── helpers for the rest ─────────────────────────
DASH_IMG = os.path.join(DECK_IMG, "dashboard.jpg")  # drop the finished dashboard screenshot here and rebuild


def placeholder(slide, x, y, w, label, hint, url):
    """A browser frame with a dashed, clearly labelled hole where a screenshot goes (a normal shape, so it is easy to select and replace)."""
    bar = 0.36
    ih = (w - 0.08) * 9 / 16
    total = bar + ih + 0.06
    shape(slide, x, y, w, total, fill=WHITE, line=RULE, line_w=1.5, radius=0.14, shadow=True)
    for i, c in enumerate(("FF5F57", "FEBC2E", "28C840")):
        shape(slide, x + 0.2 + i * 0.22, y + 0.12, 0.13, 0.13, fill=c, kind=MSO_SHAPE.OVAL)
    pill(slide, x + w * 0.28, y + 0.07, w * 0.44, 0.23, url, fill="F1EEF7", color=GREY, size=12, bold=False)
    shape(slide, x + 0.04, y + bar, w - 0.08, ih, fill=VIOLET_W, line=VIOLET, line_w=2.5, dash=True, radius=0.04, name="PLACEHOLDER dashboard screenshot")
    tb(slide, x + 0.3, y + bar + ih * 0.28, w - 0.6, 0.6, [[(label, {"bold": True})]], size=26, color=VIOLET, align="c")
    tb(slide, x + 0.3, y + bar + ih * 0.28 + 0.65, w - 0.6, 0.9, hint, size=18, color=PURPLE, align="c")
    return total


def screenshot(slide, path, x, y, w, url, ph_label=None, ph_hint=None):
    if path and os.path.exists(path):
        return browser(slide, path, x, y, w, url=url)
    return placeholder(slide, x, y, w, ph_label, ph_hint, url)


# ───────────────────────── 2. problem ─────────────────────────
set_title(s_prob, "PROBLEM STATEMENT")
claim2(s_prob, "Milestone escrow breaks in three ways,", "and most projects fix only one.")
cards = [
    (1, "Ambiguous agreements", "Agreements are written in prose. Vague terms hide disputes before any money moves."),
    (2, "Swapped deliverables", "Approval isn't tied to a specific file, so what was reviewed can differ from what gets paid."),
    (3, "Buggy money code", "A double release, a missing check or a re-entrant call can drain funds that can't be recovered."),
]
CW, CH, CY, GP = 5.3, 4.75, 3.15, 0.35
X0 = 0.9
for i, (n, title, body) in enumerate(cards):
    x = X0 + i * (CW + GP)
    shape(s_prob, x, CY, CW, CH, fill=WHITE, line=RULE, shadow=True)
    badge(s_prob, x + 0.3, CY + 0.3, 0.75, str(n), fill=NAVY)
    tb(s_prob, x + 1.25, CY + 0.28, CW - 1.5, 0.8, [[(title, {"bold": True})]], size=30, color=NAVY, anchor="m")
    tb(s_prob, x + 0.3, CY + 1.3, CW - 0.6, 1.8, body, size=24, color=INK)
my = CY + 3.25
pill(s_prob, X0 + 0.3, my, CW - 0.6, 0.6, '"to the mentor\'s satisfaction"', fill=AMBER_W, color=AMBER, size=22, bold=False)
pill(s_prob, X0 + 0.3, my + 0.78, CW - 0.6, 0.6, '"paid from the remaining funds"', fill=AMBER_W, color=AMBER, size=22, bold=False)
x2 = X0 + CW + GP
seal(s_prob, "v1.png", x2 + 0.6, my - 0.1, 1.05)
tb(s_prob, x2 + 1.7, my, 1.9, 0.95, [[("≠", {"bold": True, "font": "Segoe UI Symbol"})]], size=66, color=RED, align="c", anchor="m")
seal(s_prob, "v2.png", x2 + 3.65, my - 0.1, 1.05)
tb(s_prob, x2 + 0.15, my + 1.0, 1.8, 0.4, "reviewed", size=20, color=PURPLE, align="c")
tb(s_prob, x2 + 3.35, my + 1.0, 1.8, 0.4, "delivered", size=20, color=PURPLE, align="c")
x3 = X0 + 2 * (CW + GP)
tb(s_prob, x3 + 0.3, my - 0.15, CW - 0.6, 1.0, [[("3.000 ETH", {"bold": True})]], size=62, color=RED)
tb(s_prob, x3 + 0.3, my + 0.85, CW - 0.6, 0.6, "taken from a naive escrow in one attack", size=22, color=PURPLE)
shape(s_prob, 0.9, 8.3, 16.6, 1.9, fill=VIOLET_W, line=None, radius=0.2)
pill(s_prob, 1.25, 8.75, 3.3, 0.7, "THE BRIEF SAYS", fill=NAVY, color=WHITE, size=24)
tb(s_prob, 4.9, 8.3, 12.3, 1.9, "AI may propose a checklist, but a person must confirm it, ambiguities must be resolved, and AI output must never be written to the contract directly.", size=26, color=INK, anchor="m")
notes(s_prob, "0:20 (40 s). Three failures. One: agreements are prose, so 'to the mentor's satisfaction' is a future dispute. Two: approval is not bound to a file, so what is reviewed can differ from what is paid. Three: money code is unforgiving; our Attack Lab shows a naive escrow losing 3 ETH of other people's money to one re-entrancy attack. The booklet adds the rule that shapes our design: AI proposes, a person confirms, and AI never writes to the contract.")

# ───────────────────────── 3. idea / proposed solution ─────────────────────────
set_title(s_idea, "IDEA TITLE / PROPOSED SOLUTION")
claim2(s_idea, "Sealstone: three layers of integrity,", "so nobody can trick the payment.")
layers = [
    ("1", "AGREEMENT INTEGRITY", "AI proposes and flags unclear terms. Rules re-check, a person resolves every flag, and both sides sign the hash.", AMBER_W),
    ("2", "DELIVERABLE INTEGRITY", "The provider commits the file's SHA-256. The client can approve only that exact file.", BLUE_W),
    ("3", "FUND INTEGRITY", "Role-gated states, pull payments and a re-entrancy guard. A live ledger reconciles after every action.", GREEN_W),
]
for i, (n, title, body, wash) in enumerate(layers):
    y = 3.1 + i * 2.02
    shape(s_idea, 0.9, y, 9.4, 1.9, fill=wash, line=None, radius=0.2)
    badge(s_idea, 1.15, y + 0.32, 0.85, n, fill=NAVY, size=30)
    tb(s_idea, 2.3, y + 0.15, 7.8, 0.55, [[(title, {"bold": True})]], size=28, color=NAVY)
    tb(s_idea, 2.3, y + 0.72, 7.8, 1.15, body, size=22, color=INK)
fh = browser(s_idea, os.path.join(LAND_IMG, "desktop-vault.jpg"), 10.6, 3.1, 6.9, url="sealstone.local/vault")
pill(s_idea, 10.6, 3.1 + fh + 0.3, 6.9, 0.75, "Every milestone, and a ledger that reconciles", fill=GREEN_W, color=GREEN, size=22)
shape(s_idea, 0.9, 9.3, 16.6, 1.3, fill=NAVY, line=None, radius=0.2)
tb(s_idea, 1.3, 9.3, 15.8, 1.3, [[("AI informs. Rules decide. ", {"font": SERIF, "italic": True, "size": 30, "bold": True, "color": PEACH}), ("Every money-moving decision is made in Solidity. AI output never reaches the contract.", {"color": WHITE})]], size=26, anchor="m")
notes(s_idea, "1:00 (60 s). Our idea closes all three failures. Layer one is the agreement: the AI proposes and flags ambiguity, a deterministic validator re-checks everything, a person resolves each flag, and the keccak hash of the confirmed terms is stored on-chain and signed by both sides. Layer two is the file: approval is bound to its SHA-256. Layer three is the money: a role-gated state machine with pull payments and a live reconciliation panel. The principle: AI informs, rules decide.")

# ───────────────────────── 4. product in action ─────────────────────────
set_title(s_prod, "PRODUCT IN ACTION")
claim2(s_prod, "A working product, not a mock-up:", "judges can try every step live.")
shots = [
    (os.path.join(DECK_IMG, "studio-resolve.jpg"), "1", "AI proposes and flags", "Unclear terms are highlighted. Confirm stays locked until each is resolved.", "sealstone.local/agreement", None, None),
    (DASH_IMG, "2", "One dashboard, every vault", "What's locked, owed and waiting, with the ledger reconciled.", "sealstone.local/dashboard", "DASHBOARD SCREENSHOT", "Placeholder: the final dashboard capture goes here"),
    (os.path.join(DECK_IMG, "approve-different.jpg"), "3", "A swapped file is refused", "Two different seals. The contract reverts and funds don't move.", "sealstone.local/vault", None, None),
    (os.path.join(LAND_IMG, "desktop-lab.jpg"), "4", "12 real attacks, one click", "Each is blocked, and the ledger reconciles after every one.", "sealstone.local/attack-lab", None, None),
]
SW, SG = 3.95, 0.27
for i, (path, n, cap, sub, url, phl, phh) in enumerate(shots):
    x = 0.9 + i * (SW + SG)
    h = screenshot(s_prod, path, x, 3.2, SW, url, phl, phh)
    badge(s_prod, x - 0.18, 3.2 - 0.18, 0.62, n, fill=CORAL, color=NAVY, size=26)
    tb(s_prod, x, 3.2 + h + 0.2, SW, 0.55, [[(cap, {"bold": True})]], size=26, color=NAVY)
    tb(s_prod, x, 3.2 + h + 0.8, SW, 1.3, sub, size=21, color=INK)
tb(s_prod, 0.9, 8.05, 9.0, 0.6, [[("LIVE DEMO PATH", {"bold": True}), ("   about 5 minutes", {"color": PURPLE})]], size=28, color=NAVY)
steps = ["Extract", "Resolve flags", "Confirm + sign", "Submit file", "Swap it: refused", "Run 12 attacks"]
BW, BG = 2.55, 0.26
for i, st in enumerate(steps):
    x = 0.9 + i * (BW + BG)
    shape(s_prod, x, 8.8, BW, 1.45, fill=WHITE, line=RULE, shadow=True, radius=0.16)
    badge(s_prod, x + 0.2, 8.98, 0.55, str(i + 1), fill=NAVY, size=22)
    tb(s_prod, x + 0.2, 9.62, BW - 0.3, 0.6, [[(st, {"bold": True})]], size=21, color=NAVY)
    if i < len(steps) - 1:
        arrow(s_prod, x + BW + 0.02, 9.52, x + BW + BG - 0.02, 9.52, color=PURPLE, w=3)
notes(s_prod, "2:00 (40 s). This is the live path we will walk if you invite us to. Pasting an agreement shows the AI's flags; the dashboard shows every vault with its ledger reconciled; the swapped-file approval is refused with two visibly different seals; the Attack Lab fires twelve real transactions at the contract. Everything runs on a local chain, with Wi-Fi off if you like.")

# ───────────────────────── 5. technical approach 1/3: the pipeline ─────────────────────────
set_title(s_t1, "TECHNICAL APPROACH 1/3")
claim2(s_t1, "From a paragraph to a protected payment:", "AI proposes, rules decide, people confirm.", size=36)
NW, NH, GAP = 3.05, 2.1, 0.34
KIND = {"ai": (AMBER_W, AMBER), "rule": (VIOLET_W, VIOLET), "human": (BLUE_W, BLUE), "chain": (GREEN_W, GREEN), "plain": (WHITE, NAVY)}
nodes = [
    ("1", "Agreement text", "pasted prose, any format", "plain"),
    ("2", "AI proposes", "milestones and ambiguities (Ollama, Gemini or rules)", "ai"),
    ("3", "Rules re-check", "wei arithmetic, parties, quotes found in the text", "rule"),
    ("4", "A person resolves", "edit, remove or acknowledge every flag", "human"),
    ("5", "Confirmed checklist", "the only version that can be deployed", "chain"),
    ("6", "Seal on-chain", "keccak256 hash stored, both parties sign", "chain"),
    ("7", "Fund exactly", "the client deposits the exact total", "chain"),
    ("8", "Deliver", "the provider commits the file's SHA-256", "chain"),
    ("9", "Approve exact hash", "or dispute: the arbiter decides", "chain"),
    ("10", "Pull and reconcile", "withdraw, then the ledger checks itself", "chain"),
]
R1Y, R2Y = 3.45, 6.75
tb(s_t1, 0.9, 2.95, 9, 0.45, [[("OFF-CHAIN  ", {"bold": True}), ("AI proposes, rules and people decide", {})]], size=22, color=PURPLE)
tb(s_t1, 0.9, 6.25, 9, 0.45, [[("ON-CHAIN  ", {"bold": True}), ("Solidity decides", {})]], size=22, color=PURPLE)
pos = {}
for i, (n, title, sub, kind) in enumerate(nodes):
    col = i if i < 5 else 9 - i
    x = 0.9 + col * (NW + GAP)
    y = R1Y if i < 5 else R2Y
    pos[i] = (x, y)
    fill, ink = KIND[kind]
    shape(s_t1, x, y, NW, NH, fill=fill, line=ink, line_w=2, shadow=True, radius=0.18)
    badge(s_t1, x + 0.16, y + 0.18, 0.58, n, fill=ink, size=22)
    tb(s_t1, x + 0.85, y + 0.12, NW - 0.95, 0.7, [[(title, {"bold": True})]], size=22, color=ink, anchor="m")
    tb(s_t1, x + 0.2, y + 0.95, NW - 0.4, 1.1, sub, size=19, color=INK)
for i in range(9):
    (x1, y1), (x2, y2) = pos[i], pos[i + 1]
    if i < 4:
        arrow(s_t1, x1 + NW + 0.02, y1 + NH / 2, x2 - 0.02, y2 + NH / 2, color=NAVY)
    elif i == 4:
        arrow(s_t1, x1 + NW / 2, y1 + NH + 0.02, x2 + NW / 2, y2 - 0.02, color=CORAL, w=5)
    else:
        arrow(s_t1, x1 - 0.02, y1 + NH / 2, x2 + NW + 0.02, y2 + NH / 2, color=NAVY)
pill(s_t1, 8.0, 5.7, 7.3, 0.6, "AI BOUNDARY: only the confirmed, hashed version crosses", fill=CORAL, color=NAVY, size=20)
tb(s_t1, 0.9, 9.15, 5, 0.45, [[("COLOUR KEY", {"bold": True})]], size=20, color=PURPLE)
for i, (label_, kind) in enumerate([("AI-proposed", "ai"), ("Rules", "rule"), ("Person", "human"), ("On-chain", "chain")]):
    f, c = KIND[kind]
    pill(s_t1, 0.9 + i * 2.75, 9.7, 2.55, 0.62, label_, fill=f, color=c, size=20, line=c)
tb(s_t1, 12.1, 9.65, 5.4, 0.8, "Every refusal is a named error, and a refused call changes nothing.", size=20, color=INK)
notes(s_t1, "2:40 (50 s). Left to right, then back along the bottom. The AI proposes a checklist and ambiguities; a deterministic validator re-checks arithmetic and parties and verifies every quote exists in the text; a person resolves each flag; only the confirmed, hashed checklist crosses the AI boundary to the chain. On-chain, both parties sign the hash, the client funds the exact total, the provider commits the file's SHA-256, the client approves that exact hash or disputes, and money moves by pull payment while the ledger reconciles.")

# ───────────────────────── 6. technical approach 2/3: the state machine ─────────────────────────
set_title(s_t2, "TECHNICAL APPROACH 2/3")
claim2(s_t2, "One state machine per milestone.", "Anything not on the map reverts with a named error.", size=36)
SW_, SH_ = 2.9, 1.4


def state(slide, x, y, name, sub, fill, ink):
    shape(slide, x, y, SW_, SH_, fill=fill, line=ink, line_w=2.5, shadow=True, radius=0.2, name="State " + name)
    tb(slide, x, y + 0.12, SW_, 0.65, [[(name, {"bold": True})]], size=28, color=ink, align="c")
    tb(slide, x + 0.1, y + 0.82, SW_ - 0.2, 0.5, sub, size=19, color=INK, align="c")


def label(slide, cx, cy, w, head, sub, h=0.95, code=17):
    shape(slide, cx - w / 2, cy - h / 2, w, h, fill=WHITE, line=RULE, shadow=True, radius=0.12)
    tb(slide, cx - w / 2 + 0.1, cy - h / 2 + 0.08, w - 0.2, 0.42, [[(head, {"bold": True, "font": CODE})]], size=code, color=NAVY, align="c")
    tb(slide, cx - w / 2 + 0.1, cy - h / 2 + 0.5, w - 0.2, 0.4, sub, size=18, color=PURPLE, align="c")


P, C_, R_, D_, F_ = (0.9, 3.95), (7.2, 3.95), (13.5, 3.95), (7.2, 7.4), (13.5, 7.4)
state(s_t2, *P, "PENDING", "work not yet submitted", GREY_W, GREY)
state(s_t2, *C_, "COMPLETED", "awaiting the client", AMBER_W, AMBER)
state(s_t2, *R_, "RELEASED", "final: pays the provider", GREEN_W, GREEN)
state(s_t2, *D_, "DISPUTED", "the arbiter decides", RED_W, RED)
state(s_t2, *F_, "REFUNDED", "final: credits the client", BLUE_W, BLUE)
ym = P[1] + SH_ / 2
arrow(s_t2, P[0] + SW_ + 0.03, ym, C_[0] - 0.03, ym, color=NAVY, w=4)
label(s_t2, 5.5, 3.45, 3.2, "markComplete(i, hash)", "Provider", h=0.9)
arrow(s_t2, C_[0] + SW_ + 0.03, ym, R_[0] - 0.03, ym, color=GREEN, w=4)
label(s_t2, 11.8, 3.45, 3.2, "approve(i, hash)", "Client, the hash must match", h=0.9)
arrow(s_t2, C_[0] + SW_ / 2, C_[1] + SH_ + 0.03, D_[0] + SW_ / 2, D_[1] - 0.03, color=RED, w=4)
label(s_t2, 6.2, 6.4, 4.1, "dispute(i)  |  escalate(i)", "Client, or Provider after the window", h=0.95)
arrow(s_t2, D_[0] + SW_ + 0.03, D_[1] + 0.25, R_[0] + 1.2, R_[1] + SH_ + 0.03, color=GREEN, w=4)
label(s_t2, 12.3, 6.55, 3.0, "resolve(Release)", "Arbiter", h=0.9)
arrow(s_t2, D_[0] + SW_ + 0.03, D_[1] + SH_ / 2 + 0.2, F_[0] - 0.03, F_[1] + SH_ / 2 + 0.2, color=BLUE, w=4)
label(s_t2, 11.8, 9.0, 3.0, "resolve(Refund)", "Arbiter", h=0.9)
tb(s_t2, 0.9, 7.35, 5.6, 1.8, [[("Check order in every function", {"bold": True, "color": NAVY})], "role, agreement state, index, then milestone state. The first rule that fails names the error."], size=20, color=PURPLE, after=4)
shape(s_t2, 0.9, 9.65, 16.6, 1.0, fill=NAVY, line=None, radius=0.2)
tb(s_t2, 1.3, 9.65, 15.8, 1.0, [[("Final states never change. ", {"font": SERIF, "italic": True, "size": 30, "bold": True, "color": PEACH}), ("A second release, an unauthorized refund or a skipped step reverts with a named error.", {"color": WHITE})]], size=24, anchor="m")
notes(s_t2, "3:30 (40 s). One state machine per milestone. The provider marks it complete with the file's hash. From Completed, only the client can approve with the same hash, which credits the provider, or dispute; the provider can escalate only after the review window. Only the arbiter can settle a dispute, either releasing to the provider or refunding the client. Released and Refunded are final. Anything else, such as a second release, reverts with a specific error and changes nothing.")

# ───────────────────────── 7. technical approach 3/3: architecture ─────────────────────────
set_title(s_t3, "TECHNICAL APPROACH 3/3")
claim2(s_t3, "Three small parts,", "and the AI service can never touch the chain.")


def panel(slide, x, y, w, h, title, lines, fill=WHITE, ink=NAVY, tw=None):
    shape(slide, x, y, w, h, fill=fill, line=ink, line_w=2, shadow=True, radius=0.2)
    tb(slide, x + 0.3, y + 0.22, w - 0.6, 0.6, [[(title, {"bold": True})]], size=26, color=ink)
    tb(slide, x + 0.3, y + 0.95, tw or (w - 0.6), h - 1.05, lines, size=21, color=INK, after=6)


panel(s_t3, 0.9, 3.2, 5.9, 4.7, "WEB APP", ["Agreement studio", "Vault dashboard", "Attack Lab", [("React, TypeScript, ethers v6", {"color": PURPLE})]], tw=3.2)
s_t3.shapes.add_picture(os.path.join(LAND_IMG, "phone-milestone-mismatch.jpg"), I(4.45), I(4.05), height=I(3.55))
shape(s_t3, 4.45, 4.05, 3.55 * 780 / 1688, 3.55, fill=None, line=NAVY, line_w=2, radius=0.08)
panel(s_t3, 7.9, 3.2, 4.8, 2.55, "AI SERVICE", ["extract, validate, gate", [("Ollama, Gemini or built-in rules", {"color": PURPLE})]], fill=AMBER_W, ink=AMBER)
panel(s_t3, 7.9, 6.05, 4.8, 1.85, "SHARED PACKAGE", [[("One implementation of hashing, validator and gate", {"color": INK})]], fill=VIOLET_W, ink=VIOLET)
panel(s_t3, 13.9, 3.2, 3.6, 4.7, "LOCAL CHAIN", ["MilestoneVault.sol", "OpenZeppelin guard", [("Attack contracts, test only", {"color": PURPLE})], [("Solidity 0.8.24, Hardhat", {"color": PURPLE})]], fill=GREEN_W, ink=GREEN)
arrow(s_t3, 6.83, 4.4, 7.88, 4.4, color=NAVY, w=4)
tb(s_t3, 6.8, 4.55, 1.1, 0.5, "HTTP", size=18, color=PURPLE, align="c")
arrow(s_t3, 3.0, 7.92, 3.0, 8.45, color=NAVY, w=4, head=False)
arrow(s_t3, 3.0, 8.45, 15.7, 8.45, color=NAVY, w=4, head=False)
arrow(s_t3, 15.7, 8.45, 15.7, 7.94, color=NAVY, w=4)
pill(s_t3, 6.95, 8.15, 4.9, 0.6, "JSON-RPC: sign, fund, approve, withdraw", fill=CREAM, color=NAVY, size=20)
arrow(s_t3, 12.72, 4.4, 13.88, 4.4, color=RED, w=4, dash=True, head=False)
badge(s_t3, 12.95, 4.05, 0.7, "✕", fill=RED, size=26)
tb(s_t3, 12.65, 4.85, 1.25, 0.9, ["never", "connects"], size=18, color=RED, align="c")
shape(s_t3, 0.9, 9.1, 16.6, 1.5, fill=NAVY, line=None, radius=0.2)
tb(s_t3, 1.3, 9.1, 15.8, 1.5, [[("Local fallback. ", {"font": SERIF, "italic": True, "size": 30, "bold": True, "color": PEACH}), ("With Wi-Fi off it still works: rules extractor, local chain, bundled fonts. No paid API, GPU or cloud.", {"color": WHITE})]], size=26, anchor="m")
notes(s_t3, "4:10 (30 s). Three parts. The web app, which also works on a phone, talks to the chain directly over JSON-RPC. The AI service extracts, validates and gates, and it has no route, key or RPC connection to the chain: it only stores confirmed JSON. One shared package implements hashing, the validator and the gate, so the browser, server and deploy scripts agree byte for byte. Everything is local, so the full demo works with Wi-Fi off.")

# ───────────────────────── 8. feasibility and viability ─────────────────────────
set_title(s_feas, "FEASIBILITY AND VIABILITY")
claim2(s_feas, "Built, tested and reproducible in 12 hours,", "with no paid service.")
tb(s_feas, 0.9, 3.0, 4.4, 0.45, [[("BOOKLET RULE", {"bold": True})]], size=22, color=PURPLE)
tb(s_feas, 5.75, 3.0, 5.0, 0.45, [[("EVIDENCE", {"bold": True})]], size=22, color=PURPLE)
rows = [
    ("Working MVP, testable live", "npm run dev: studio, vault, Attack Lab"),
    ("Reproducible live testing", "demo:reset, a seeded vault, 383 tests"),
    ("AI used meaningfully", "extracts milestones, flags ambiguity"),
    ("Checkable rules decide", "validator and Solidity; AI never overrides"),
    ("Fictional data only", "invented agreements, test ETH, local chain"),
    ("Local fallback", "rules extractor and local chain, Wi-Fi off"),
]
for i, (rule, ev) in enumerate(rows):
    y = 3.55 + i * 1.05
    shape(s_feas, 0.9, y, 9.9, 0.92, fill=WHITE, line=RULE, shadow=False, radius=0.14)
    tb(s_feas, 1.1, y, 0.5, 0.92, [[("✓", {"bold": True, "font": "Segoe UI Symbol"})]], size=26, color=GREEN, anchor="m")
    tb(s_feas, 1.75, y, 3.9, 0.92, [[(rule, {"bold": True})]], size=21, color=NAVY, anchor="m")
    tb(s_feas, 5.75, y, 4.95, 0.92, ev, size=20, color=INK, anchor="m")
stats = [("138", "contract tests passing"), ("100%", "line coverage of the escrow"), ("1,000", "fuzz steps, 8 invariants"), ("12/12", "real attacks blocked")]
for i, (big, small) in enumerate(stats):
    x = 11.1 + (i % 2) * 3.3
    y = 3.0 + (i // 2) * 2.75
    shape(s_feas, x, y, 3.1, 2.55, fill=NAVY, line=None, radius=0.2, shadow=True)
    tb(s_feas, x, y + 0.2, 3.1, 1.3, [[(big, {"bold": True})]], size=60, color=PEACH, align="c", anchor="m")
    tb(s_feas, x + 0.2, y + 1.55, 2.7, 0.9, small, size=20, color=WHITE, align="c")
shape(s_feas, 11.1, 8.6, 6.4, 1.95, fill=CREAM, line=NAVY, line_w=2, radius=0.14)
tb(s_feas, 11.4, 8.6, 5.9, 1.95, [[("$ ", {"color": PURPLE, "font": CODE}), ("npm run dev", {"font": CODE})], [("$ ", {"color": PURPLE, "font": CODE}), ("npm run test:contracts", {"font": CODE})], [("$ ", {"color": PURPLE, "font": CODE}), ("npm run demo:reset", {"font": CODE})]], size=19, color=NAVY, anchor="m", after=4)
notes(s_feas, "4:40 (35 s). Feasibility is evidenced, not claimed. The booklet's Common Rules map to concrete proof: a working product judges can run, reproducible resets, meaningful AI with deterministic rules deciding, fictional data, and a local fallback. The numbers are from our own runs: 138 contract tests, 383 tests overall, 100% line coverage, a 1,000-step randomized invariant fuzz, and twelve real attacks all blocked.")

# ───────────────────────── 9. impact and benefits ─────────────────────────
set_title(s_imp, "IMPACT AND BENEFITS")
claim2(s_imp, "Fewer disputes, no swapped files,", "and no hidden intermediary to trust.")
shape(s_imp, 0.9, 3.1, 9.4, 5.9, fill=WHITE, line=RULE, shadow=True, radius=0.2)
tb(s_imp, 1.25, 3.25, 8.7, 0.6, [[("Other users' money left in the contract", {"bold": True})]], size=24, color=NAVY)
cd = CategoryChartData()
cd.categories = ["0", "1", "2", "3", "4"]
cd.add_series("Sealstone", (3, 3, 3, 3, 3))
cd.add_series("Naive escrow", (3, 3, 2, 1, 0))
gf = s_imp.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, I(1.1), I(3.95), I(9.0), I(4.95), cd)
gf.name = "Chart: money left, naive vs Sealstone"
ch = gf.chart
ch.font.name, ch.font.size = HR, Pt(20)
ch.font.color.rgb = rgb(INK)
ch.has_title = False
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(22)
ch.legend.font.name = HR
va = ch.value_axis
va.minimum_scale, va.maximum_scale, va.major_unit = 0, 3.6, 1
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = rgb(RULE)
va.format.line.fill.background()
va.tick_labels.number_format = '0" ETH"'
va.tick_labels.number_format_is_linked = False
va.tick_labels.font.size = Pt(22)
va.tick_labels.font.name = HR
ca = ch.category_axis
ca.has_title = True
ca.axis_title.text_frame.text = "Payouts made to the attacker"
tr = ca.axis_title.text_frame.paragraphs[0].runs[0]
tr.font.size, tr.font.bold, tr.font.name = Pt(22), False, HR
ca.tick_labels.font.size = Pt(22)
ca.tick_labels.font.name = HR
ca.format.line.color.rgb = rgb(RULE)
for ser, colr, last in ((ch.plots[0].series[0], VIOLET, "3 ETH safe"), (ch.plots[0].series[1], RED, "0 ETH left")):
    ser.smooth = False
    ser.format.line.color.rgb = rgb(colr)
    ser.format.line.width = Pt(5)
    ser.marker.style = XL_MARKER_STYLE.CIRCLE
    ser.marker.size = 12
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = rgb(WHITE)
    ser.marker.format.line.color.rgb = rgb(colr)
    ser.marker.format.line.width = Pt(3)
    dl = ser.points[4].data_label
    dl.has_text_frame = True
    dl.text_frame.text = last
    r = dl.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.name = Pt(24), True, HB
    r.font.color.rgb = rgb(colr)
    dl.position = XL_LABEL_POSITION.ABOVE if colr == VIOLET else XL_LABEL_POSITION.LEFT
tb(s_imp, 0.9, 9.1, 9.4, 0.8, "One re-entrancy attack, payout by payout. In the Attack Lab's contrast run the attacker takes 4 ETH and 3 ETH of other users' money is lost.", size=19, color=PURPLE)
who = [
    ("Freelancers and clients", "Pay only for the exact file you approved. Disputes go to a named referee."),
    ("Grant programmes", "One dashboard shows what's locked, owed and waiting across every agreement."),
    ("Auditors and finance", "The ledger reconciles after every action, and the rules are tests anyone can run."),
]
for i, (title, body) in enumerate(who):
    y = 3.1 + i * 2.0
    shape(s_imp, 10.6, y, 6.9, 1.88, fill=WHITE, line=RULE, shadow=True, radius=0.18)
    badge(s_imp, 10.85, y + 0.3, 0.72, str(i + 1), fill=CORAL, color=NAVY, size=26)
    tb(s_imp, 11.8, y + 0.15, 5.55, 0.6, [[(title, {"bold": True})]], size=24, color=NAVY)
    tb(s_imp, 11.8, y + 0.72, 5.5, 1.1, body, size=20, color=INK)
tb(s_imp, 0.9, 9.95, 3.0, 0.6, [[("WHAT'S NEXT", {"bold": True})]], size=22, color=PURPLE, anchor="m")
for i, t in enumerate(["Factory for many agreements", "A committee as arbiter", "Pilot on a public testnet"]):
    pill(s_imp, 4.0 + i * 4.5, 9.95, 4.3, 0.65, t, fill=VIOLET_W, color=VIOLET, size=20)
notes(s_imp, "5:15 (35 s). The impact is visible in one picture. The same re-entrancy attack drains a naive escrow of 3 ETH of other people's money, and cannot touch ours. That translates into three groups of people: freelancers and clients who pay only for the exact file they approved, grant programmes that see everything locked and owed in one dashboard, and auditors who can reproduce every guarantee as a test. Next steps are a factory for many agreements, a committee arbiter, and a public-testnet pilot.")

# ───────────────────────── 10. differentiator ─────────────────────────
set_title(s_diff, "DIFFERENTIATOR")
claim2(s_diff, "Most escrow demos secure the money.", "Sealstone also secures the agreement and the file.", size=38)
cards10 = [
    ("AGREEMENT INTEGRITY", AMBER, AMBER_W, "Vague terms can't reach the contract", "agreement studio, Extract"),
    ("DELIVERABLE INTEGRITY", BLUE, BLUE_W, "Approval is bound to the exact file", "rejects approval of a swapped deliverable"),
    ("FUND INTEGRITY", GREEN, GREEN_W, "Forced-in money can't break the books", "attack: Force ETH in with selfdestruct"),
    ("ADVERSARIAL PROOF", VIOLET, VIOLET_W, "You fire the attacks yourself", "Attack Lab, Run all attacks"),
]
CW10, GAP10, CY10, CH10 = 3.95, 0.27, 3.1, 5.4
for i, (eyebrow, ink, wash, head, proof) in enumerate(cards10):
    x = 0.9 + i * (CW10 + GAP10)
    shape(s_diff, x, CY10, CW10, CH10, fill=WHITE, line=ink, line_w=2.5, shadow=True, radius=0.2)
    pill(s_diff, x + 0.25, CY10 + 0.25, CW10 - 0.5, 0.5, eyebrow, fill=wash, color=ink, size=17)
    tb(s_diff, x + 0.3, CY10 + 0.95, CW10 - 0.6, 1.2, [[(head, {"bold": True})]], size=26, color=NAVY)
    shape(s_diff, x + 0.25, CY10 + 4.15, CW10 - 0.5, 1.1, fill=CREAM, line=None, radius=0.12)
    tb(s_diff, x + 0.4, CY10 + 4.2, CW10 - 0.8, 0.3, [[("LIVE PROOF", {"bold": True})]], size=14, color=PURPLE)
    tb(s_diff, x + 0.4, CY10 + 4.5, CW10 - 0.8, 0.7, [[(proof, {"font": CODE})]], size=15, color=NAVY)
vy = CY10 + 2.3
# 1: a vague phrase, resolved by a person
x = 0.9
pill(s_diff, x + 0.35, vy, CW10 - 0.7, 0.55, '"to the mentor\'s satisfaction"', fill=AMBER_W, color=AMBER, size=19, bold=False)
tb(s_diff, x + 0.35, vy + 0.55, CW10 - 0.7, 0.45, [[("↓", {"bold": True, "font": "Segoe UI Symbol"})]], size=26, color=PURPLE, align="c", anchor="m")
pill(s_diff, x + 0.35, vy + 1.05, CW10 - 0.7, 0.55, "A live demo, set by a person", fill=GREEN_W, color=GREEN, size=19, bold=False)
# 2: two seals that differ, and the named refusal
x = 0.9 + (CW10 + GAP10)
seal(s_diff, "v1.png", x + 0.45, vy - 0.05, 0.95)
tb(s_diff, x + 1.4, vy, 1.2, 0.85, [[("≠", {"bold": True, "font": "Segoe UI Symbol"})]], size=48, color=RED, align="c", anchor="m")
seal(s_diff, "v2.png", x + 2.55, vy - 0.05, 0.95)
shape(s_diff, x + 0.3, vy + 1.05, CW10 - 0.6, 0.55, fill=RED_W, line=None, radius=0.27)
tb(s_diff, x + 0.3, vy + 1.05, CW10 - 0.6, 0.55, [[("DeliverableHashMismatch", {"font": CODE, "bold": True})]], size=15, color=RED, align="c", anchor="m")
# 3: forced ETH shown as surplus
x = 0.9 + 2 * (CW10 + GAP10)
tb(s_diff, x + 0.3, vy - 0.15, CW10 - 0.6, 0.9, [[("0.300 ETH", {"bold": True})]], size=44, color=VIOLET, align="c", anchor="m")
tb(s_diff, x + 0.3, vy + 0.7, CW10 - 0.6, 0.4, "shown as surplus, never withdrawable", size=17, color=PURPLE, align="c")
pill(s_diff, x + 0.5, vy + 1.15, CW10 - 1.0, 0.5, "✓  books still reconcile", fill=GREEN_W, color=GREEN, size=19)
# 4: twelve real attacks
x = 0.9 + 3 * (CW10 + GAP10)
tb(s_diff, x + 0.3, vy - 0.15, CW10 - 0.6, 0.9, [[("12 / 12", {"bold": True})]], size=48, color=NAVY, align="c", anchor="m")
tb(s_diff, x + 0.3, vy + 0.7, CW10 - 0.6, 0.4, "real transactions, all blocked", size=17, color=PURPLE, align="c")
pill(s_diff, x + 0.5, vy + 1.15, CW10 - 1.0, 0.5, "Run all attacks", fill=NAVY, color=WHITE, size=19)
# the closing challenge
shape(s_diff, 0.9, 8.75, 16.6, 1.15, fill=NAVY, line=None, radius=0.2)
tb(s_diff, 1.3, 8.75, 15.8, 1.15, [[("Don't take our word for it. ", {"font": SERIF, "italic": True, "bold": True, "color": PEACH, "size": 30}), ("Pick a role in the demo vault, or press Run all attacks.", {"color": WHITE, "size": 26})]], size=26, anchor="m")
pill(s_diff, 0.9, 10.15, 4.4, 0.55, "AI never moves money", fill=VIOLET_W, color=VIOLET, size=18)
pill(s_diff, 5.5, 10.15, 3.9, 0.55, "Works with Wi-Fi off", fill=VIOLET_W, color=VIOLET, size=18)
tb(s_diff, 9.7, 10.05, 7.8, 0.8, [[("Honest limits: ", {"bold": True, "color": NAVY}), ("one contract per agreement, a trusted arbiter, and a hash proves identity, not quality.", {})]], size=17, color=PURPLE, anchor="m")
notes(s_diff, "5:50 (30 s). Each of these is something you can check live, not a claim. One: a vague phrase like 'to the mentor's satisfaction' cannot reach the contract until a person resolves it. Two: approval is bound to the exact file, so a look-alike is refused with a named error, DeliverableHashMismatch. Three: if someone force-sends ETH into the vault, it is counted as surplus, nobody can withdraw it, and the books still reconcile. Four: you can fire twelve real attacks at the contract yourself and read the result. The AI never moves money, the demo works with Wi-Fi off, and we state our limits: one contract per agreement, a trusted arbiter, and a hash proves identity, not quality. Don't take our word for it: pick a role, or press Run all attacks.")

# ───────────────────────── 11. thank you ─────────────────────────
seal(s_thanks, "brand.png", 9.35, 0.55, 1.3)
tb(s_thanks, 3.6, 1.95, 12.8, 0.65, [[("SEALSTONE  ", {"bold": True}), ("escrow that can't be tricked", {})]], size=30, color=PURPLE, align="c")
tb(s_thanks, 3.6, 2.6, 12.8, 0.6, "Team TORVALDS  ·  XQR1520246  ·  Omkar Rane (leader), Pranay Patil", size=24, color=NAVY, align="c")
notes(s_thanks, "6:20. Thank you. Everything we showed is running on this laptop, offline. We would love you to try to break it: pick a role in the demo vault, or press Run all in the Attack Lab. Team TORVALDS, team ID XQR1520246: Omkar Rane and Pranay Patil.")

prs.save(OUT)
print("saved", OUT)
