"""Building blocks for the Sealstone pitch deck (python-pptx). Slide canvas: 20 x 11.25 in (the Vortexa template)."""
import math
import os
from lxml import etree
from PIL import Image, ImageDraw
from pptx.chart.data import CategoryChartData  # noqa: F401  (used by build.py)
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ── the template's own typefaces (embedded in the file) and palette ──
HB = "Heading Now 31-38 Bold"
HR = "Heading Now 31-38"
CODE = "Consolas"
SERIF = "Georgia"  # accent face for one-line takeaways (ships with Office on Windows and Mac); never used for numbers
NAVY, INK, PURPLE = "172243", "161F3E", "4E4A76"
CORAL, PEACH, LAV, CREAM = "EE7A6B", "FAD3AE", "D6C2DB", "FDF9ED"
WHITE, RULE = "FFFFFF", "E6DCE4"
# the app's state colours, used ONLY with their meaning: amber = AI-proposed, blue = edited by a person, green = on-chain/confirmed, red = refused
AMBER, AMBER_W = "9A5B0A", "FBEBCF"
BLUE, BLUE_W = "2F56C9", "E0E8FB"
GREEN, GREEN_W = "0B7A5B", "DDF3EA"
RED, RED_W = "C0342B", "FBE1DE"
GREY, GREY_W = "6B6878", "ECEAF3"
VIOLET, VIOLET_W = "5B45E6", "EFEBFF"


SCALE = 1.18  # the template's typeface is condensed, so text can run larger than its point size suggests


def rgb(h):
    return RGBColor.from_string(h)


def I(v):
    return Inches(v)


def _shadow(shape, blur=16, dist=5, alpha=18, color="1E2761"):
    """A soft drop shadow. Must be called AFTER fill and line are set (schema order)."""
    sp = shape._element.spPr
    eff = sp.find(qn("a:effectLst"))  # `shadow.inherit = False` already created an (empty) one: fill it, never add a second
    if eff is None:
        eff = etree.SubElement(sp, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(blur * 12700)), dist=str(int(dist * 12700)), dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val=color)
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def shape(slide, x, y, w, h, fill=WHITE, line=None, line_w=1.5, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14, shadow=False, dash=False, name=None):
    s = slide.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            from pptx.enum.dml import MSO_LINE

            s.line.dash_style = MSO_LINE.DASH
    s.shadow.inherit = False
    if shadow:
        _shadow(s)
    if name:
        s.name = name
    s.text_frame.text = ""
    return s


def tb(slide, x, y, w, h, paras, size=24, font=HR, color=INK, align="l", anchor="t", spacing=1.0, after=0, name=None):
    """Text box. `paras`: str | [paragraph]; paragraph: str | [run]; run: str | (text, {size,font,color,bold})."""
    s = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = s.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, (str, tuple)):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        para.line_spacing = spacing
        para.space_after = Pt(after)
        runs = p if isinstance(p, list) else [p]
        for r in runs:
            text, o = (r, {}) if isinstance(r, str) else r
            run = para.add_run()
            run.text = text
            f = run.font
            f.name = o.get("font", HB if o.get("bold") else font)
            fnt = o.get("font", HB if o.get("bold") else font)
            sz = o.get("size", size)
            f.size = Pt(sz if fnt in (CODE, SERIF, "Segoe UI Symbol") else round(sz * SCALE, 1))
            f.bold = True if o.get("bold") else None
            if o.get("italic"):
                f.italic = True
            f.color.rgb = rgb(o.get("color", color))
    if name:
        s.name = name
    return s


def arrow(slide, x1, y1, x2, y2, color=NAVY, w=3.5, head=True, dash=False, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        from pptx.enum.dml import MSO_LINE

        c.line.dash_style = MSO_LINE.DASH
    if head:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="lg", len="lg")
    if name:
        c.name = name
    return c


def badge(slide, x, y, d, label, fill=NAVY, color=WHITE, size=26):
    """A numbered circle."""
    s = shape(slide, x, y, d, d, fill=fill, kind=MSO_SHAPE.OVAL)
    tb(slide, x, y, d, d, label, size=size, font=HB, color=color, align="c", anchor="m")
    return s


def pill(slide, x, y, w, h, label, fill=VIOLET_W, color=VIOLET, size=22, bold=True, line=None):
    s = shape(slide, x, y, w, h, fill=fill, line=line, radius=h / 2)
    tb(slide, x, y, w, h, [[(label, {"bold": bold})]], size=size, color=color, align="c", anchor="m")
    return s


def seal_png(path, hexstr, color="172243", px=640):
    """The app's hash seal (one radial stroke per hex digit), drawn the same way as the HashSeal component."""
    hexstr = hexstr.replace("0x", "").ljust(64, "0")
    S = px * 2
    k = S / 100.0
    img = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    col = tuple(int(color[i : i + 2], 16) for i in (0, 2, 4)) + (255,)
    for i in range(64):
        v = int(hexstr[i], 16)
        a = (i / 64) * math.pi * 2 - math.pi / 2
        r1, r2 = 22, 26 + v * 1.25
        d.line([(50 * k + r1 * k * math.cos(a), 50 * k + r1 * k * math.sin(a)), (50 * k + r2 * k * math.cos(a), 50 * k + r2 * k * math.sin(a))], fill=col, width=int(1.9 * k))
    d.ellipse([50 * k - 17 * k, 50 * k - 17 * k, 50 * k + 17 * k, 50 * k + 17 * k], outline=col, width=int(1.3 * k))
    for j in range(10):  # dashed inner ring
        a0 = j * 36
        d.arc([50 * k - 8 * k, 50 * k - 8 * k, 50 * k + 8 * k, 50 * k + 8 * k], a0, a0 + 20, fill=col, width=int(1.3 * k))
    d.ellipse([50 * k - 2.4 * k, 50 * k - 2.4 * k, 50 * k + 2.4 * k, 50 * k + 2.4 * k], fill=col)
    img = img.resize((px, px), Image.LANCZOS)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    img.save(path)
    return path


def browser(slide, img_path, x, y, w, url="sealstone.local", radius=0.14):
    """A framed screenshot (16:9). Returns the total height used."""
    bar = 0.36
    with Image.open(img_path) as im:
        ratio = im.height / im.width
    ih = (w - 0.08) * ratio
    total = bar + ih + 0.06
    shape(slide, x, y, w, total, fill=WHITE, line=RULE, line_w=1.5, radius=radius, shadow=True)
    for i, c in enumerate(("FF5F57", "FEBC2E", "28C840")):
        shape(slide, x + 0.2 + i * 0.22, y + 0.12, 0.13, 0.13, fill=c, kind=MSO_SHAPE.OVAL)
    pill(slide, x + w * 0.28, y + 0.07, w * 0.44, 0.23, url, fill="F1EEF7", color=GREY, size=12, bold=False)
    slide.shapes.add_picture(img_path, I(x + 0.04), I(y + bar), width=I(w - 0.08))
    return total


def set_title(slide, text):
    """The template's section title: left-aligned, nudged down (its top clipped), brush scaled to the words."""
    t = next(s for s in slide.shapes if s.name == "TextBox 10")
    t.text_frame.paragraphs[0].runs[0].text = text
    t.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    t.left, t.top, t.width = I(0.37), I(0.16), I(16.0)
    brush = next(s for s in slide.shapes if s.name == "Freeform 11")
    brush.top = I(1.30)
    brush.width = I(max(5.49, 0.343 * len(text) + 0.3))


def claim(slide, text, size=40):
    return tb(slide, 0.85, 2.0, 16.4, 1.0, [[(text, {"bold": True})]], size=size, color=NAVY, name="Claim")


def claim2(slide, head, tail, size=40):
    """Two-tone takeaway: a bold condensed head, then a serif-italic tail in the template's purple."""
    return tb(slide, 0.85, 2.0, 16.6, 1.0, [[(head + " ", {"bold": True, "color": NAVY, "size": size}), (tail, {"font": SERIF, "italic": True, "color": PURPLE, "size": size * 0.72})]], size=size, color=NAVY, anchor="m", name="Claim")


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
